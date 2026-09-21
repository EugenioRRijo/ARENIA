"""El camino de la pantalla llega al mismo sitio que el del catalogo (Tarea 3, compuerta GP1).

Los 1 586,61 USD y el 7 de 7 (indicador 1 de la tesis) ya estan demostrados **elaborando desde el
catalogo**: `tests/integration/test_presupuesto_linea_base.py` y
`tests/integration/test_auditoria_7_de_7.py`. Esta prueba no los vuelve a demostrar por ese camino
ni redeclara ningun dato: lo unico que aporta es que **el camino nuevo** -- el que recorre una
persona cuando teclea las tres tablas de `ui/paginas/componer.py` y estas pasan por
`ui.composicion.composicion_desde_tablas` -- produce exactamente las mismas `ComposicionAPU`, el
mismo `Presupuesto` y el mismo informe de auditoria. Si las dos rutas divergieran, el prototipo
mentiria.

De donde sale cada valor esperado
=================================

De ningun sitio de este archivo. La linea base (los cinco APU linea a linea, la curva emitida, el
total auditado y las siete inconsistencias con su regla) vive **una sola vez** en
`tests/fixtures/apu_linea_base.py`, y el presupuesto con las siete inconsistencias en
`tests/fixtures/presupuesto_auditado.py` (CLAUDE.md §2, principio DRY). Aqui se importan; no se
transcribe ni un precio, ni una cantidad, ni una descripcion. Una prueba que copiara sus valores
esperados dejaria de ser una regresion en cuanto alguien corrigiera la fixture: seguiria afirmando
el numero viejo y pasaria mientras el sistema se desvia.

El caso de demostracion tampoco se construye aqui: se importa de `scripts/seed_demo.py`
(`construir_caso_demo`), su unica construccion en el repositorio. **Ese caso es didactico y no
proviene de obra ejecutada** (CLAUDE.md §1): nadie tendio esa tuberia ni instalo esa valvula; sus
precios son de ejemplo, no cotizados. Existe para ejercitar en un mismo presupuesto lo que la
linea base no tiene: jornal y destajo en la misma partida (decision D9, articulo 114 de la LOTTT),
materiales con desperdicio, equipos con depreciacion parcial, la curva con `plan_secuencial` y la
exportacion a Excel con la columna Modalidad que anadio la Tarea 1.

Ninguna prueba de este archivo se omite: la compuerta GP1 no admite un *skip* silencioso donde
deberia estar la regresion del indicador 1.
"""

from __future__ import annotations

from decimal import Decimal

import pytest
from openpyxl import load_workbook

from core.budget import con_curva, elaborar, exportar_excel, generar_presupuesto, plan_secuencial
from core.budget.excel import HOJA_APU, HOJA_AUDITORIA, HOJA_CURVA, HOJA_PRESUPUESTO
from core.contracts import ComposicionAPU, ModalidadManoObra, Presupuesto
from core.verification import auditar
from core.verification.informe import DECIMALES_PRESENTACION
from scripts.seed_demo import CasoDemo, construir_caso_demo
from tests.fixtures import apu_linea_base as linea_base
from tests.fixtures.computo_auditado import composiciones_linea_base
from tests.fixtures.presupuesto_auditado import (
    items_con_trazas,
    presupuesto_con_siete_inconsistencias,
)
from tests.integration.test_auditoria_7_de_7 import _detectadas
from ui.composicion import composicion_desde_tablas

#: Tolerancia del total contra `TOTAL_PRESUPUESTO_AUDITADO`: un presupuesto se lee en centavos y
#: los `Decimal` del modelo no se redondean hasta presentarlos (CLAUDE.md §2.3 y §4). No se
#: teclea "0.01": se deriva de la precision de presentacion que ya declara el nucleo, la misma de
#: la que `core/budget/excel.py` deduce su paso de redondeo.
TOLERANCIA = Decimal(1).scaleb(-DECIMALES_PRESENTACION)

#: Las inconsistencias de la linea base que el sistema debe detectar. El numero es la afirmacion
#: medible de la tesis (indicador 1), no un dato copiado de una fixture: por eso se escribe, y por
#: eso se comprueba ademas contra `INCONSISTENCIAS`, que es donde viven las siete de verdad.
INCONSISTENCIAS_DE_LA_TESIS = 7

#: Una unidad de material por unidad de partida. Por encima de este umbral, un material medido en
#: la misma unidad que su partida lleva desperdicio (convencion de `tests/fixtures/apu_linea_base`
#: y de `scripts/seed_demo.py`: 1,05 = 1 m de tuberia + 5 % de merma).
UNIDAD_COMPLETA = Decimal(1)

#: Encabezados de la tabla de mano de obra de la hoja APU (`core/budget/excel.py`). Son el objeto
#: de la comprobacion, no valores esperados copiados de una fixture, y ese modulo los escribe como
#: literales, sin constante que importar.
ENCABEZADO_MANO_OBRA = "Mano de obra"
ENCABEZADO_MODALIDAD = "Modalidad"


# --------------------------------------------------------------------------------------------
# El camino de la pantalla: de un `ComposicionAPU` a las tres tablas de texto y de vuelta
# --------------------------------------------------------------------------------------------


def _texto(valor: Decimal) -> str:
    """El `Decimal` tal como lo teclea una persona: notacion fija, sin exponente.

    `str(Decimal("1E+2"))` da `"1E+2"`, que ninguna pantalla muestra; `f"{...:f}"` fuerza la
    notacion posicional sin perder ni un digito ni la escala (`Decimal("1.00")` -> `"1.00"`).
    """
    return f"{valor:f}"


def _filas_materiales(apu: ComposicionAPU) -> list[dict[str, str]]:
    return [
        {
            "descripcion": linea.descripcion,
            "unidad": linea.unidad,
            "cantidad": _texto(linea.cantidad),
            "precio": _texto(linea.precio),
        }
        for linea in apu.materiales
    ]


def _filas_equipos(apu: ComposicionAPU) -> list[dict[str, str]]:
    return [
        {
            "descripcion": equipo.descripcion,
            "cantidad": _texto(equipo.cantidad),
            "precio": _texto(equipo.precio),
            "depreciacion": _texto(equipo.depreciacion),
        }
        for equipo in apu.equipos
    ]


def _filas_mano_obra(apu: ComposicionAPU) -> list[dict[str, str]]:
    return [
        {
            "descripcion": obrero.descripcion,
            "cantidad": _texto(obrero.cantidad),
            "sueldo": _texto(obrero.sueldo),
            "modalidad": obrero.modalidad.value,
        }
        for obrero in apu.mano_obra
    ]


def _tecleado(apu: ComposicionAPU) -> ComposicionAPU:
    """El mismo APU rehecho por el camino de la pantalla, sin tocar `core`.

    Desarma la composicion en las tres tablas de diccionarios **de texto** que produce
    `st.data_editor` y se las da a `ui.composicion.composicion_desde_tablas`, que es exactamente
    lo que hace `ui/paginas/componer.py` al guardar. Lo que vuelve debe ser igual al original: la
    igualdad de dataclass compara codigo, descripcion, unidad, rendimiento y las tres tuplas de
    lineas campo a campo, igual que `tests/integration/test_persistencia.py` compara la ida y
    vuelta por SQLite.
    """
    return composicion_desde_tablas(
        apu.codigo_partida,
        apu.descripcion,
        apu.unidad,
        _texto(apu.rendimiento),
        _filas_materiales(apu),
        _filas_equipos(apu),
        _filas_mano_obra(apu),
    )


# --------------------------------------------------------------------------------------------
# Regresion de la linea base: 1 586,61 USD y 7 de 7, por el camino de la pantalla
# --------------------------------------------------------------------------------------------


def _presupuesto_tecleado() -> Presupuesto:
    """El presupuesto auditado, con las cinco composiciones rehechas por la pantalla.

    Las cantidades, las trazas y la curva son las de la fixture (`items_con_trazas` y
    `CURVA_AUDITADA`): lo unico que cambia frente a `presupuesto_con_siete_inconsistencias` es por
    donde pasaron los APU. La curva emitida se coloca con `con_curva` y no con `plan_secuencial`
    porque es precisamente la que **no** cierra (inconsistencia 3): un plan generado por el nucleo
    cerraria al 100 % por construccion y R2 se quedaria callada.
    """
    borrador = generar_presupuesto(
        items_con_trazas(),
        {codigo: _tecleado(apu) for codigo, apu in composiciones_linea_base().items()},
        linea_base.PARAMETROS_LINEA_BASE,
        codigo=linea_base.CODIGO_PRESUPUESTO,
        fecha=linea_base.FECHA_LINEA_BASE,
        moneda=linea_base.MONEDA,
    )
    return con_curva(borrador, linea_base.CURVA_AUDITADA)


@pytest.mark.parametrize(
    "apu", linea_base.APUS_LINEA_BASE, ids=lambda apu: apu.codigo_partida
)
def test_la_composicion_tecleada_es_la_de_la_linea_base(apu: ComposicionAPU) -> None:
    """Cada uno de los cinco APU, tecleado en la pantalla, es el APU de la fixture.

    Si esta prueba falla, las dos rutas divergen y el hallazgo es de la investigacion: el mensaje
    de pytest senala el campo y la linea exactos.
    """
    assert _tecleado(apu) == apu


def test_el_presupuesto_tecleado_es_el_del_catalogo() -> None:
    """La igualdad completa: mismas partidas, mismos `ResultadoAPU`, misma curva."""
    assert _presupuesto_tecleado() == presupuesto_con_siete_inconsistencias()


def test_el_presupuesto_tecleado_reproduce_el_total_auditado() -> None:
    """El total del caso, por el camino de la pantalla (importado, nunca escrito)."""
    presupuesto = _presupuesto_tecleado()

    diferencia = abs(presupuesto.total - linea_base.TOTAL_PRESUPUESTO_AUDITADO)

    assert diferencia <= TOLERANCIA, (
        f"total tecleado {presupuesto.total} contra el auditado "
        f"{linea_base.TOTAL_PRESUPUESTO_AUDITADO}"
    )


def test_la_auditoria_del_presupuesto_tecleado_detecta_las_siete() -> None:
    """Indicador 1 de la tesis por el camino de la pantalla: 7 de 7.

    Las siete esperadas no se enumeran aqui: son las que declara `INCONSISTENCIAS` en la fixture
    de la linea base, la misma declaracion que usa `tests/integration/test_auditoria_7_de_7.py`.
    """
    informe = auditar(_presupuesto_tecleado())

    esperadas = {
        (inconsistencia.regla, inconsistencia.codigo_partida)
        for inconsistencia in linea_base.INCONSISTENCIAS
    }
    detectadas = _detectadas(informe)

    assert len(esperadas) == INCONSISTENCIAS_DE_LA_TESIS
    assert esperadas <= detectadas, f"no detectadas: {sorted(esperadas - detectadas, key=str)}"
    assert len(esperadas & detectadas) == INCONSISTENCIAS_DE_LA_TESIS


def test_la_auditoria_tecleada_es_la_del_catalogo() -> None:
    """Y no solo las siete: el informe entero coincide, hallazgo por hallazgo."""
    assert auditar(_presupuesto_tecleado()) == auditar(presupuesto_con_siete_inconsistencias())


# --------------------------------------------------------------------------------------------
# El caso de demostracion (didactico, NO proviene de obra ejecutada)
# --------------------------------------------------------------------------------------------


def _elaborado_demo(caso: CasoDemo):
    """El presupuesto del caso de demostracion, con sus APU rehechos por la pantalla y su curva.

    Caso didactico: sus tres partidas son un ejercicio academico ficticio y sus precios son de
    ejemplo, no cotizados (docstring de `scripts/seed_demo.py`, CLAUDE.md §1).
    """
    composiciones = {apu.codigo_partida: _tecleado(apu) for apu in caso.composiciones}
    borrador = generar_presupuesto(
        caso.items,
        composiciones,
        caso.parametros,
        codigo=caso.codigo_presupuesto,
        fecha=caso.fecha,
        moneda=caso.moneda,
    )
    return elaborar(
        caso.items,
        composiciones,
        caso.parametros,
        codigo=caso.codigo_presupuesto,
        fecha=caso.fecha,
        moneda=caso.moneda,
        plan=plan_secuencial(borrador),
    )


def test_el_caso_de_demostracion_sobrevive_al_camino_de_la_pantalla() -> None:
    """Las tres composiciones del caso didactico, tecleadas, son las mismas.

    Importa en particular la linea a destajo: `modalidad` viaja como texto por la tabla de mano de
    obra y debe volver como `ModalidadManoObra.DESTAJO`, no degradarse al JORNAL por defecto.
    """
    for apu in construir_caso_demo().composiciones:
        assert _tecleado(apu) == apu


def test_el_caso_de_demostracion_mezcla_jornal_y_destajo_en_una_partida() -> None:
    """Decision D9 (articulo 114 de la LOTTT) ejercitada de punta a punta.

    Las dos modalidades no se escriben: se toman de `ModalidadManoObra`, que es quien las declara.
    """
    modalidades = {
        apu.codigo_partida: {linea.modalidad for linea in _tecleado(apu).mano_obra}
        for apu in construir_caso_demo().composiciones
    }

    mixtas = [codigo for codigo, usadas in modalidades.items() if usadas == set(ModalidadManoObra)]

    assert mixtas, f"ninguna partida del caso mezcla las dos modalidades: {modalidades}"


def test_el_caso_de_demostracion_lleva_desperdicio_y_depreciacion_parcial() -> None:
    """Material medido en la unidad de su partida y consumido a mas de uno por unidad (merma), y
    equipo imputado por una fraccion de su precio (depreciacion parcial).

    Ambas condiciones se escriben como propiedades, no como los valores concretos del caso: si
    manana `seed_demo` cambia el 1,05 por 1,07, la prueba sigue siendo valida.
    """
    composiciones = [_tecleado(apu) for apu in construir_caso_demo().composiciones]

    con_desperdicio = [
        (apu.codigo_partida, linea.descripcion)
        for apu in composiciones
        for linea in apu.materiales
        if linea.unidad == apu.unidad and linea.cantidad > UNIDAD_COMPLETA
    ]
    con_depreciacion_parcial = [
        (apu.codigo_partida, equipo.descripcion)
        for apu in composiciones
        for equipo in apu.equipos
        if equipo.depreciacion < UNIDAD_COMPLETA
    ]

    assert con_desperdicio, "ningun material del caso lleva desperdicio"
    assert con_depreciacion_parcial, "ningun equipo del caso se imputa por depreciacion parcial"


def test_la_curva_del_caso_de_demostracion_cierra_en_el_total() -> None:
    """Lo que la linea base no hacia (inconsistencia 3): con `plan_secuencial`, la curva del
    nucleo cierra exactamente en el total y R2 no tiene nada que senalar."""
    resultado = _elaborado_demo(construir_caso_demo())

    assert resultado.presupuesto.curva
    assert resultado.presupuesto.total_curva == resultado.presupuesto.total


def test_exportar_excel_escribe_las_cuatro_hojas_con_la_modalidad(tmp_path) -> None:
    """Las cuatro hojas del libro, y la columna Modalidad de la Tarea 1 con su dato real.

    No basta con que el encabezado exista: se lee la columna en cada bloque de APU y se compara
    con las modalidades de la composicion que lo produjo. Un encabezado sobre una columna vacia
    reproduciria justo el malentendido que la decision D9 existe para evitar.
    """
    caso = construir_caso_demo()
    resultado = _elaborado_demo(caso)

    ruta = exportar_excel(resultado.presupuesto, resultado.informe, tmp_path / "demostracion.xlsx")

    libro = load_workbook(ruta)
    assert set(libro.sheetnames) == {HOJA_PRESUPUESTO, HOJA_APU, HOJA_CURVA, HOJA_AUDITORIA}

    filas = list(libro[HOJA_APU].iter_rows(values_only=True))
    encabezados = [
        indice
        for indice, fila in enumerate(filas)
        if fila[0] == ENCABEZADO_MANO_OBRA and ENCABEZADO_MODALIDAD in fila
    ]
    assert len(encabezados) == len(resultado.presupuesto.partidas)

    for indice, partida in zip(encabezados, resultado.presupuesto.partidas, strict=True):
        columna = filas[indice].index(ENCABEZADO_MODALIDAD)
        lineas = partida.apu.mano_obra
        leidas = [fila[columna] for fila in filas[indice + 1 : indice + 1 + len(lineas)]]
        assert leidas == [linea.modalidad.value for linea in lineas], partida.apu.codigo_partida
