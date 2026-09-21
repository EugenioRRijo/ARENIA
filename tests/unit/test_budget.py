"""Presupuesto, curva de inversión y exportación a Excel (Sesión I0.5).

El presupuesto se arma con las cantidades del caso auditado y los cinco APU de la línea base, así
que su total debe reproducir los 1 586,61 USD del PDF con la tolerancia de ± 0,01 de CLAUDE.md §4.

La curva es el punto crítico de la sesión: en el presupuesto del caso didáctico cerraba en 99,30 %
(hallazgo 3), y aquí debe cerrar exactamente en el total **por construcción**, incluso repartiendo
con montos redondeados a dos decimales.

Ningún dato se repite: cantidades, APU y totales vienen de `tests.fixtures.apu_linea_base` a través
de `tests.fixtures.computo_auditado` (principio DRY de CLAUDE.md §2).
"""

from __future__ import annotations

from datetime import date
from decimal import Decimal

import pytest
from openpyxl import load_workbook

from core.budget import (
    con_curva,
    elaborar,
    exportar_excel,
    generar_curva,
    generar_presupuesto,
    plan_secuencial,
)
from core.budget.curva import PlanInvalido
from core.contracts.apu import ComposicionAPU, LineaManoObra, ModalidadManoObra, ParametrosCosto
from core.contracts.dominio import Dominio
from core.contracts.item_computo import ItemComputo, OrigenTipo
from core.verification.directivas import codigos_en_etiqueta
from core.verification.informe import InformeAuditoria
from tests.fixtures import apu_linea_base as linea_base
from tests.fixtures.computo_auditado import composiciones_linea_base, items_auditados

TOLERANCIA = Decimal("0.01")
CENTIMO = Decimal("0.01")
COMPLETO = Decimal("1")
MITAD = Decimal("0.5")

HOJAS_ESPERADAS = ["Presupuesto", "APU", "Curva", "Auditoria"]


def _presupuesto():
    return generar_presupuesto(
        items_auditados(),
        composiciones_linea_base(),
        linea_base.PARAMETROS_LINEA_BASE,
        codigo=linea_base.CODIGO_PRESUPUESTO,
        fecha=linea_base.FECHA_LINEA_BASE,
        moneda=linea_base.MONEDA,
    )


def _plan_auditado():
    """Los seis períodos del plan de trabajo del caso: el encofrado se ejecuta en dos mitades."""
    return [
        ("Dia 1", ((linea_base.APU_EXCAVACION.codigo_partida, COMPLETO),)),
        ("Dia 2", ((linea_base.APU_TUBERIA.codigo_partida, COMPLETO),)),
        ("Dia 3", ((linea_base.APU_ENCOFRADO.codigo_partida, MITAD),)),
        ("Dia 4", ((linea_base.APU_ENCOFRADO.codigo_partida, MITAD),)),
        ("Dia 5", ((linea_base.APU_CONCRETO.codigo_partida, COMPLETO),)),
        ("Dia 6", ((linea_base.APU_RELLENO.codigo_partida, COMPLETO),)),
    ]


def _acumulados_coherentes(curva) -> bool:
    corrido = Decimal(0)
    for punto in curva:
        corrido += punto.monto
        if punto.acumulado != corrido:
            return False
    return True


def test_presupuesto_reproduce_1586_61():
    presupuesto = _presupuesto()

    assert len(presupuesto.partidas) == len(linea_base.PRESUPUESTO_AUDITADO)
    assert abs(presupuesto.total - linea_base.TOTAL_PRESUPUESTO_AUDITADO) <= TOLERANCIA


def test_cada_renglon_reproduce_el_total_auditado():
    presupuesto = _presupuesto()

    for partida, linea in zip(presupuesto.partidas, linea_base.PRESUPUESTO_AUDITADO, strict=True):
        assert partida.item.codigo_partida == linea.codigo_partida
        assert abs(partida.resultado.precio_unitario - linea.precio_unitario) <= TOLERANCIA
        assert abs(partida.total - linea.total) <= TOLERANCIA, linea.codigo_partida


def test_codigo_sin_composicion_lanza_valueerror():
    composiciones = composiciones_linea_base()
    huerfano = linea_base.APU_RELLENO.codigo_partida
    del composiciones[huerfano]

    with pytest.raises(ValueError, match=huerfano):
        generar_presupuesto(
            items_auditados(),
            composiciones,
            linea_base.PARAMETROS_LINEA_BASE,
            codigo=linea_base.CODIGO_PRESUPUESTO,
            fecha=linea_base.FECHA_LINEA_BASE,
        )


def test_curva_cierra_exactamente_en_el_total():
    presupuesto = _presupuesto()

    curva = generar_curva(presupuesto, _plan_auditado(), cuantizar=CENTIMO)
    cerrado = con_curva(presupuesto, curva)

    assert len(curva) == len(linea_base.CURVA_AUDITADA)
    assert _acumulados_coherentes(curva)
    assert cerrado.total_curva == cerrado.total
    assert presupuesto.curva == (), "con_curva no debe modificar el presupuesto original"
    assert codigos_en_etiqueta(curva[0].periodo) == (linea_base.APU_EXCAVACION.codigo_partida,)


@pytest.mark.parametrize(
    "fracciones",
    [
        (Decimal("1"),),
        (Decimal("0.5"), Decimal("0.5")),
        (Decimal("0.25"), Decimal("0.25"), Decimal("0.5")),
        (Decimal("0.125"), Decimal("0.375"), Decimal("0.2"), Decimal("0.3")),
    ],
)
def test_curva_con_fracciones_arbitrarias_cierra(fracciones):
    presupuesto = _presupuesto()
    codigos = [partida.item.codigo_partida for partida in presupuesto.partidas]
    plan = [
        (f"Etapa {numero}", tuple((codigo, fraccion) for codigo in codigos))
        for numero, fraccion in enumerate(fracciones, start=1)
    ]

    curva = generar_curva(presupuesto, plan, cuantizar=CENTIMO)

    assert len(curva) == len(fracciones)
    assert _acumulados_coherentes(curva)
    assert curva[-1].acumulado == presupuesto.total


def test_fracciones_que_no_suman_uno_lanzan_valueerror():
    presupuesto = _presupuesto()
    incompleto = linea_base.APU_ENCOFRADO.codigo_partida
    plan = [
        (
            etiqueta,
            tuple(
                (codigo, MITAD if codigo == incompleto else fraccion)
                for codigo, fraccion in reparto
            ),
        )
        for etiqueta, reparto in plan_secuencial(presupuesto)
    ]

    with pytest.raises(PlanInvalido, match=incompleto) as error:
        generar_curva(presupuesto, plan)
    assert isinstance(error.value, ValueError)


def test_codigo_desconocido_en_el_plan_lanza_valueerror():
    presupuesto = _presupuesto()
    plan = [*plan_secuencial(presupuesto), ("Dia 6", (("LB-99-XXX", COMPLETO),))]

    with pytest.raises(PlanInvalido, match="LB-99-XXX"):
        generar_curva(presupuesto, plan)


def test_elaborar_siempre_devuelve_informe():
    sin_plan = elaborar(
        items_auditados(),
        composiciones_linea_base(),
        linea_base.PARAMETROS_LINEA_BASE,
        codigo=linea_base.CODIGO_PRESUPUESTO,
        fecha=linea_base.FECHA_LINEA_BASE,
    )
    con_plan = elaborar(
        items_auditados(),
        composiciones_linea_base(),
        linea_base.PARAMETROS_LINEA_BASE,
        codigo=linea_base.CODIGO_PRESUPUESTO,
        fecha=linea_base.FECHA_LINEA_BASE,
        plan=_plan_auditado(),
    )

    assert isinstance(sin_plan.informe, InformeAuditoria)
    assert isinstance(con_plan.informe, InformeAuditoria)
    assert sin_plan.presupuesto.curva == ()
    assert con_plan.presupuesto.total_curva == con_plan.presupuesto.total
    assert con_plan.informe.codigo_presupuesto == linea_base.CODIGO_PRESUPUESTO


def test_exportar_excel_escribe_cuatro_hojas(tmp_path):
    resultado = elaborar(
        items_auditados(),
        composiciones_linea_base(),
        linea_base.PARAMETROS_LINEA_BASE,
        codigo=linea_base.CODIGO_PRESUPUESTO,
        fecha=linea_base.FECHA_LINEA_BASE,
        plan=_plan_auditado(),
    )
    ruta = tmp_path / "presupuesto.xlsx"

    escrita = exportar_excel(resultado.presupuesto, resultado.informe, ruta)

    assert escrita == ruta
    libro = load_workbook(ruta)
    assert libro.sheetnames == HOJAS_ESPERADAS

    hoja = libro["Presupuesto"]
    total_escrito = Decimal(str(hoja.cell(row=hoja.max_row, column=hoja.max_column).value))
    assert abs(total_escrito - linea_base.TOTAL_PRESUPUESTO_AUDITADO) <= TOLERANCIA
    assert libro["Curva"].max_row == len(resultado.presupuesto.curva) + 1
    assert libro["Auditoria"].max_row == len(resultado.informe.hallazgos) + 1


def test_hoja_apu_muestra_la_modalidad_de_mano_de_obra(tmp_path):
    """La columna Modalidad evita el malentendido de la decision D9: bajo destajo, `sueldo` es el
    precio por unidad de partida, no un jornal diario, y quien audite el libro debe poder verlo."""
    item = ItemComputo(
        codigo_partida="X-01",
        descripcion="Partida con mano de obra mixta",
        unidad="m2",
        cantidad=Decimal("1"),
        origen_id="TEST-MODALIDAD-01",
        origen_tipo=OrigenTipo.MANUAL,
        dominio=Dominio.CIVIL,
    )
    apu = ComposicionAPU(
        codigo_partida="X-01",
        descripcion="Partida con mano de obra mixta",
        unidad="m2",
        rendimiento=Decimal("50"),
        mano_obra=(
            LineaManoObra("Ayudante a jornal", Decimal("2"), Decimal("3")),
            LineaManoObra(
                "Instalador a destajo",
                Decimal("1"),
                Decimal("6"),
                ModalidadManoObra.DESTAJO,
            ),
        ),
    )
    resultado = elaborar(
        [item],
        {"X-01": apu},
        ParametrosCosto(),
        codigo="PRES-TEST-MODALIDAD",
        fecha=date(2026, 9, 20),
    )
    ruta = tmp_path / "modalidad.xlsx"

    exportar_excel(resultado.presupuesto, resultado.informe, ruta)

    hoja = load_workbook(ruta)["APU"]
    filas = list(hoja.iter_rows(values_only=True))
    indice_encabezado = next(i for i, fila in enumerate(filas) if fila[0] == "Mano de obra")
    encabezado = filas[indice_encabezado]
    fila_jornal = filas[indice_encabezado + 1]
    fila_destajo = filas[indice_encabezado + 2]

    assert "Modalidad" in encabezado
    assert encabezado[-1] == "Total"
    assert fila_jornal[-2] == "jornal"
    assert fila_destajo[-2] == "destajo"
    # El Total de esta tabla es cantidad x sueldo (igual que en materiales y equipos), no el
    # aporte con FCAS y bono que va en el resumen: aqui coincide en 6.00 para ambas lineas
    # (2 obreros x 3 = 1 destajista x 6), pero por razones distintas segun la modalidad.
    assert fila_jornal[-1] == Decimal("6.00")
    assert fila_destajo[-1] == Decimal("6.00")
