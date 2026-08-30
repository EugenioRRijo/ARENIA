"""Exportación del presupuesto a Excel: presupuesto, APU, curva de inversión y auditoría.

Es el único lugar del núcleo donde se redondea: las celdas se escriben con dos decimales porque un
presupuesto se lee en centavos, mientras que los `Decimal` del modelo conservan todos sus dígitos
(CLAUDE.md §2.3). La cantidad de decimales no se declara aquí: se toma de
`core.verification.informe.DECIMALES_PRESENTACION`, que ya la fija para el informe de auditoría.

El libro lleva siempre la hoja de auditoría, aunque el presupuesto no tenga hallazgos: el informe se
entrega siempre (principio 7 de CLAUDE.md §2).
"""

from __future__ import annotations

from decimal import ROUND_HALF_UP, Decimal
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.worksheet.worksheet import Worksheet

from core.contracts.presupuesto import PartidaPresupuestada, Presupuesto
from core.verification.informe import DECIMALES_PRESENTACION, InformeAuditoria

__all__ = ["exportar_excel"]

HOJA_PRESUPUESTO = "Presupuesto"
HOJA_APU = "APU"
HOJA_CURVA = "Curva"
HOJA_AUDITORIA = "Auditoria"

PASO_REDONDEO = Decimal(1).scaleb(-DECIMALES_PRESENTACION)
CIEN = Decimal(100)
CERO = Decimal(0)

_NEGRITA = Font(bold=True)


def exportar_excel(presupuesto: Presupuesto, informe: InformeAuditoria, ruta: Path) -> Path:
    """Escribe el libro con las cuatro hojas y devuelve la ruta escrita."""
    libro = Workbook()
    _hoja_presupuesto(libro.active, presupuesto)
    _hoja_apu(libro.create_sheet(HOJA_APU), presupuesto)
    _hoja_curva(libro.create_sheet(HOJA_CURVA), presupuesto)
    _hoja_auditoria(libro.create_sheet(HOJA_AUDITORIA), informe)
    libro.save(ruta)
    return ruta


def _redondear(valor: Decimal) -> Decimal:
    return valor.quantize(PASO_REDONDEO, rounding=ROUND_HALF_UP)


def _encabezado(hoja: Worksheet, titulos: list[str]) -> None:
    hoja.append(titulos)
    for celda in hoja[hoja.max_row]:
        celda.font = _NEGRITA


def _hoja_presupuesto(hoja: Worksheet, presupuesto: Presupuesto) -> None:
    hoja.title = HOJA_PRESUPUESTO
    _encabezado(
        hoja,
        ["N.º", "Código", "Descripción", "Unidad", "Cantidad", "Precio unitario", "Total"],
    )
    for numero, partida in enumerate(presupuesto.partidas, start=1):
        hoja.append(
            [
                numero,
                partida.item.codigo_partida,
                partida.item.descripcion,
                partida.item.unidad,
                _redondear(partida.item.cantidad),
                _redondear(partida.resultado.precio_unitario),
                _redondear(partida.total),
            ]
        )
    fila = hoja.max_row + 1
    hoja.cell(row=fila, column=6, value="TOTAL").font = _NEGRITA
    hoja.cell(row=fila, column=7, value=_redondear(presupuesto.total)).font = _NEGRITA


def _hoja_apu(hoja: Worksheet, presupuesto: Presupuesto) -> None:
    for partida in presupuesto.partidas:
        _bloque_apu(hoja, partida)
        hoja.append([])


def _bloque_apu(hoja: Worksheet, partida: PartidaPresupuestada) -> None:
    apu = partida.apu
    hoja.append([f"{apu.codigo_partida} · {apu.descripcion}"])
    hoja[hoja.max_row][0].font = _NEGRITA
    hoja.append(["Unidad", apu.unidad, "Rendimiento", _redondear(apu.rendimiento)])

    _encabezado(hoja, ["Materiales", "Unidad", "Cantidad", "Precio", "Total"])
    for linea in apu.materiales:
        hoja.append(
            [
                linea.descripcion,
                linea.unidad,
                _redondear(linea.cantidad),
                _redondear(linea.precio),
                _redondear(linea.total),
            ]
        )

    _encabezado(hoja, ["Equipos", "Cantidad", "Precio", "Depreciación", "Total"])
    for equipo in apu.equipos:
        hoja.append(
            [
                equipo.descripcion,
                _redondear(equipo.cantidad),
                _redondear(equipo.precio),
                equipo.depreciacion,
                _redondear(equipo.total),
            ]
        )

    _encabezado(hoja, ["Mano de obra", "Obreros", "Sueldo", "Total"])
    for obrero in apu.mano_obra:
        hoja.append(
            [
                obrero.descripcion,
                _redondear(obrero.cantidad),
                _redondear(obrero.sueldo),
                _redondear(obrero.total),
            ]
        )

    resultado = partida.resultado
    for etiqueta, valor in (
        ("Materiales", resultado.materiales),
        ("Equipos", resultado.equipos),
        ("Mano de obra", resultado.mano_obra),
        ("Costo directo", resultado.costo_directo),
        ("Con administración", resultado.con_administracion),
        ("Precio unitario", resultado.precio_unitario),
    ):
        hoja.append([etiqueta, _redondear(valor)])


def _hoja_curva(hoja: Worksheet, presupuesto: Presupuesto) -> None:
    _encabezado(hoja, ["Período", "Monto", "Acumulado", "% acumulado"])
    total = presupuesto.total
    for punto in presupuesto.curva:
        porcentaje = CERO if total == CERO else punto.acumulado / total * CIEN
        hoja.append(
            [
                punto.periodo,
                _redondear(punto.monto),
                _redondear(punto.acumulado),
                _redondear(porcentaje),
            ]
        )


def _hoja_auditoria(hoja: Worksheet, informe: InformeAuditoria) -> None:
    _encabezado(hoja, ["Regla", "Severidad", "Descripción", "Impacto", "Orígenes"])
    for hallazgo in informe.hallazgos:
        hoja.append(
            [
                hallazgo.regla,
                hallazgo.severidad.name,
                hallazgo.descripcion,
                None if hallazgo.impacto is None else _redondear(hallazgo.impacto),
                ", ".join(hallazgo.origen_ids),
            ]
        )
